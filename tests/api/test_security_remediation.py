from __future__ import annotations

import asyncio
import csv
from datetime import UTC, datetime, timedelta
from decimal import Decimal
from io import BytesIO, StringIO
from types import SimpleNamespace
from uuid import UUID

import pyotp
import pytest
from helpers import (
    authenticate,
    create_activated_user,
    office_id,
    owner_client,
    spawned_client,
)
from nexa_bos_api.applications.models import (
    Application,
    ApplicationEvent,
    WorkflowStage,
)
from nexa_bos_api.case_operations.csv_service import (
    HEADERS,
    current_cases_csv,
    validate_rows,
)
from nexa_bos_api.case_operations.reporting import REPORT_COLUMNS, case_report_xlsx
from nexa_bos_api.core.exceptions import AppError
from nexa_bos_api.core.image_storage import validate_image_header
from nexa_bos_api.core.spreadsheets import spreadsheet_safe
from nexa_bos_api.identity.access import load_user_with_type
from nexa_bos_api.identity.models import OneTimeToken, User
from nexa_bos_api.main import app
from nexa_bos_api.reporting.export import build_excel
from nexa_bos_api.reporting.periods import PeriodWindow
from nexa_bos_api.reporting.tl import (
    MAX_HISTORY_POINTS,
    _history_cutoffs,
    _metric_history,
)
from nexa_bos_api.targets import service as target_service
from openpyxl import load_workbook
from PIL import Image
from sqlalchemy import select, text
from test_applications import _catalog, _create_app, _customer, _enable_case_owner
from test_contracts import _contract_type, _create_contract, _signed_attachment
from test_leave_management import _annual_leave
from test_scope import _scoped_type


@pytest.mark.asyncio
async def test_scope_assignment_cannot_expand_self_or_move_peer_outside_office(client):
    owner, _ = await owner_client(client)
    dxb, auh = await office_id(owner, "DXB"), await office_id(owner, "AUH")
    code = await _scoped_type(owner, "office", ["Users.View", "Users.Edit"])
    actor = await create_activated_user(owner, user_type_code=code, office_id=dxb)
    peer = await create_activated_user(owner, office_id=dxb)
    foreign = await create_activated_user(owner, office_id=auh)
    async with await spawned_client() as scoped:
        await authenticate(scoped, actor["email"], "UserPass1!")
        for destination in (auh, None):
            response = await scoped.patch(
                f"/api/v1/users/{actor['id']}",
                json={"office_id": destination, "department_id": None, "team_id": None},
            )
            assert response.status_code == 403, response.text
        denied = await scoped.patch(f"/api/v1/users/{peer['id']}", json={"office_id": auh})
        assert denied.status_code == 403, denied.text
        allowed = await scoped.patch(
            f"/api/v1/users/{actor['id']}",
            json={"office_id": dxb, "full_name": "Safe self edit"},
        )
        assert allowed.status_code == 200, allowed.text
        assert (await scoped.get(f"/api/v1/users/{foreign['id']}")).status_code == 403
    assert (await owner.get(f"/api/v1/users/{actor['id']}")).json()["office"]["id"] == dxb
    assert (await owner.get(f"/api/v1/users/{peer['id']}")).json()["office"]["id"] == dxb
    # Company-authorized administration remains functional.
    moved = await owner.patch(f"/api/v1/users/{peer['id']}", json={"office_id": auh})
    assert moved.status_code == 200, moved.text


@pytest.mark.asyncio
async def test_peer_history_requires_view_and_audit_remains_separate(client):
    owner, _ = await owner_client(client)
    no_view = await _scoped_type(owner, "company", ["Users.Edit"])
    view = await _scoped_type(owner, "company", ["Users.View"])
    actor = await create_activated_user(owner, user_type_code=no_view)
    viewer = await create_activated_user(owner, user_type_code=view)
    peer = await create_activated_user(owner)
    async with await spawned_client() as caller:
        await authenticate(caller, actor["email"], "UserPass1!")
        assert (await caller.get(f"/api/v1/users/{peer['id']}/history")).status_code == 403
        own = await caller.get(f"/api/v1/users/{actor['id']}/history")
        assert own.status_code == 200, own.text
    async with await spawned_client() as caller:
        await authenticate(caller, viewer["email"], "UserPass1!")
        visible = await caller.get(f"/api/v1/users/{peer['id']}/history")
        assert visible.status_code == 200, visible.text
        assert visible.json()["events"] == []


@pytest.mark.asyncio
async def test_independent_customer_scope_also_protects_self_assignment(client):
    owner, _ = await owner_client(client)
    code = await _scoped_type(owner, "company", ["Users.View", "Users.Edit", "Customers.View"])
    user_type = next(
        item
        for item in (await owner.get("/api/v1/user-types")).json()["items"]
        if item["code"] == code
    )
    for path, field, scope in (
        ("application-scope", "application_visibility_scope", "company"),
        ("reporting-scope", "reporting_visibility_scope", "company"),
        ("customer-scope", "customer_visibility_scope", "office"),
    ):
        configured = await owner.put(
            f"/api/v1/user-types/{user_type['id']}/{path}", json={field: scope}
        )
        assert configured.status_code == 200, configured.text
    dxb, auh = await office_id(owner, "DXB"), await office_id(owner, "AUH")
    actor = await create_activated_user(owner, user_type_code=code, office_id=dxb)
    async with await spawned_client() as caller:
        await authenticate(caller, actor["email"], "UserPass1!")
        denied = await caller.patch(
            f"/api/v1/users/{actor['id']}",
            json={"office_id": auh, "full_name": "Must roll back"},
        )
        assert denied.status_code == 403, denied.text
    retained = (await owner.get(f"/api/v1/users/{actor['id']}")).json()
    assert retained["office"]["id"] == dxb
    assert retained["fullName"] == actor["fullName"]


@pytest.mark.asyncio
async def test_rejected_decode_and_storage_failure_preserve_existing_profile_photo(
    client, monkeypatch, tmp_path
):
    owner, _ = await owner_client(client)
    user = await create_activated_user(owner)
    monkeypatch.setattr("nexa_bos_api.identity.users_service.storage_dir", lambda: tmp_path)
    path = f"/api/v1/users/{user['id']}/photo"
    image = BytesIO()
    Image.new("RGB", (120, 120), "blue").save(image, "JPEG")
    original = image.getvalue()
    uploaded = await owner.post(path, files={"file": ("original.jpg", original, "image/jpeg")})
    assert uploaded.status_code == 200, uploaded.text
    before = {item.name: item.read_bytes() for item in tmp_path.iterdir()}
    corrupt = original[:-10]
    # JPEG structural verify accepts this; actual pixel decoding must reject it.
    assert validate_image_header(corrupt, "image/jpeg") == "JPEG"
    rejected = await owner.post(path, files={"file": ("bad.jpg", corrupt, "image/jpeg")})
    assert rejected.status_code == 422, rejected.text
    assert {item.name: item.read_bytes() for item in tmp_path.iterdir()} == before
    assert (await owner.get(path)).content == original
    corrupt_png = bytearray(_png(1, 1))
    idat = corrupt_png.index(b"IDAT")
    chunk_length = int.from_bytes(corrupt_png[idat - 4 : idat], "big")
    corrupt_png[idat + 4 + chunk_length] ^= 1
    rejected_png = await owner.post(
        path, files={"file": ("bad-checksum.png", bytes(corrupt_png), "image/png")}
    )
    assert rejected_png.status_code == 422, rejected_png.text
    assert rejected_png.json()["error"]["code"] == "IMAGE_CONTENT_INVALID"
    assert {item.name: item.read_bytes() for item in tmp_path.iterdir()} == before
    assert (await owner.get(path)).content == original
    from pathlib import Path

    real_write = Path.write_bytes
    writes = 0

    def interrupted_write(file, data):
        nonlocal writes
        writes += 1
        if writes == 2:
            raise OSError("Synthetic storage interruption")
        return real_write(file, data)

    with monkeypatch.context() as interrupted:
        interrupted.setattr(Path, "write_bytes", interrupted_write)
        with pytest.raises(OSError, match="Synthetic storage interruption"):
            await owner.post(path, files={"file": ("replacement.jpg", original, "image/jpeg")})
    assert {item.name: item.read_bytes() for item in tmp_path.iterdir()} == before
    assert (await owner.get(path)).content == original


def _png(width, height):
    buffer = BytesIO()
    Image.new("RGB", (width, height), "red").save(buffer, "PNG")
    return buffer.getvalue()


@pytest.mark.parametrize("dimensions", [(5000, 1), (4096, 4096)])
@pytest.mark.asyncio
async def test_profile_photo_rejects_decoded_bounds_before_write_or_transpose(
    client, monkeypatch, tmp_path, dimensions
):
    owner, _ = await owner_client(client)
    user = await create_activated_user(owner)
    monkeypatch.setattr("nexa_bos_api.identity.users_service.storage_dir", lambda: tmp_path)
    payload = _png(*dimensions)
    assert len(payload) < 2 * 1024 * 1024
    with monkeypatch.context() as decoding:

        def forbidden_transpose(*args, **kwargs):
            pytest.fail("Unbounded image reached pixel decoding")

        decoding.setattr(
            "nexa_bos_api.identity.users_service.ImageOps.exif_transpose",
            forbidden_transpose,
        )
        response = await owner.post(
            f"/api/v1/users/{user['id']}/photo",
            files={"file": ("compressed.png", payload, "image/png")},
        )
    assert response.status_code == 422, response.text
    assert response.json()["error"]["code"] == "IMAGE_DIMENSIONS_INVALID"
    assert list(tmp_path.iterdir()) == []
    assert validate_image_header(_png(96, 96), "image/png") == "PNG"
    with pytest.raises(AppError):
        validate_image_header(_png(96, 96), "image/jpeg")


@pytest.mark.parametrize(
    "value",
    ["=1+1", "+SUM(1,1)", "-1+1", "@SUM(1,1)", "  =1+1", "\tplain", "\r=1", "\n=1"],
)
def test_spreadsheet_strings_are_literal_and_numbers_preserved(value):
    safe = spreadsheet_safe(value)
    assert safe == "'" + value
    assert spreadsheet_safe(Decimal("-12.50")) == Decimal("-12.50")
    assert spreadsheet_safe(7) == 7
    key = next(key for key, _ in REPORT_COLUMNS if key == "bankFileNumber")
    book = load_workbook(BytesIO(case_report_xlsx({"items": [{key: value}]})))
    column = next(i for i, (field, _) in enumerate(REPORT_COLUMNS, 1) if field == key)
    # XML parsers normalize literal carriage returns to line feeds.
    assert book.active.cell(2, column).value == safe.replace("\r", "\n")
    assert book.active.cell(2, column).data_type == "s"
    metadata = load_workbook(
        BytesIO(
            build_excel(
                title="Security control",
                actor=SimpleNamespace(full_name=value, user_code="TEST"),
                payload={"items": []},
                filters={},
            )
        )
    )
    generated_by = next(row[1] for row in metadata["Metadata"] if row[0].value == "Generated By")
    assert generated_by.value.startswith("'")
    assert generated_by.data_type == "s"


@pytest.mark.asyncio
async def test_contract_history_and_replaced_files_require_history_permission(client):
    owner, _ = await owner_client(client)
    contract_type = await _contract_type(owner)
    hr = await create_activated_user(owner, user_type_code="HR")
    employee = await create_activated_user(owner)
    view_code = await _scoped_type(owner, "company", ["Contracts.View"])
    history_code = await _scoped_type(owner, "company", ["Contracts.View", "Contracts.History"])
    only_history_code = await _scoped_type(owner, "company", ["Contracts.History"])
    viewer = await create_activated_user(owner, user_type_code=view_code)
    historian = await create_activated_user(owner, user_type_code=history_code)
    history_only = await create_activated_user(owner, user_type_code=only_history_code)
    async with await spawned_client() as hr_client:
        await authenticate(hr_client, hr["email"], "UserPass1!")
        contract = await _create_contract(
            hr_client,
            employee_id=employee["id"],
            type_id=contract_type["id"],
            number=f"SEC-{employee['id']}",
        )
        await _signed_attachment(hr_client, contract["id"])
        reviewed = (await owner.get(f"/api/v1/contracts/{contract['id']}")).json()
        old = next(item for item in reviewed["attachments"] if item["isActive"])
        replaced = await hr_client.post(
            f"/api/v1/contracts/{contract['id']}/attachment",
            files={"upload": ("replacement.pdf", b"%PDF-1.4\n%%EOF", "application/pdf")},
            data={"reason": "New signed evidence"},
        )
        assert replaced.status_code == 200, replaced.text
        current = (await owner.get(f"/api/v1/contracts/{contract['id']}")).json()
        assert current["lockVersion"] > reviewed["lockVersion"]
        assert len(current["attachments"]) == 2
    url = f"/api/v1/contracts/{contract['id']}"
    file_url = f"{url}/attachments/{old['id']}/file"
    for user, permitted in ((viewer, False), (historian, True)):
        async with await spawned_client() as caller:
            await authenticate(caller, user["email"], "UserPass1!")
            detail = await caller.get(url)
            assert detail.status_code == 200, detail.text
            assert bool(detail.json()["history"]) is permitted
            assert all(item["isActive"] for item in detail.json()["attachments"]) is (not permitted)
            listing = await caller.get("/api/v1/contracts")
            listed = next(item for item in listing.json()["items"] if item["id"] == contract["id"])
            assert bool(listed["history"]) is permitted
            assert (await caller.get(file_url)).status_code == (200 if permitted else 404)
    async with await spawned_client() as caller:
        await authenticate(caller, history_only["email"], "UserPass1!")
        assert (await caller.get(url)).status_code == 403


@pytest.mark.asyncio
async def test_cancellation_attempt_requires_fresh_manager_consent(client):
    owner, _ = await owner_client(client)
    configured = await owner.put(
        "/api/v1/attendance/working-days", json={"weekdays": [0, 1, 2, 3, 4]}
    )
    assert configured.status_code == 200, configured.text
    annual = await _annual_leave(owner)
    manager = await create_activated_user(owner, user_type_code="TL")
    employee = await create_activated_user(owner, manager_id=manager["id"])
    hr = await create_activated_user(owner, user_type_code="HR")
    async with (
        await spawned_client() as worker,
        await spawned_client() as lead,
        await spawned_client() as human_resources,
    ):
        for caller, user in (
            (worker, employee),
            (lead, manager),
            (human_resources, hr),
        ):
            await authenticate(caller, user["email"], "UserPass1!")
        created = await worker.post(
            "/api/v1/leave/requests",
            json={
                "leave_type_id": annual["id"],
                "start_date": "2027-02-01",
                "end_date": "2027-02-03",
                "portion": "full_day",
                "reason": "Synthetic security regression",
                "submit": True,
            },
        )
        assert created.status_code == 200, created.text
        path = f"/api/v1/leave/requests/{created.json()['id']}"
        state = created.json()
        for caller, action in (
            (lead, "manager-approve"),
            (human_resources, "hr-approve"),
        ):
            response = await caller.post(
                f"{path}/{action}", json={"lock_version": state["lockVersion"]}
            )
            assert response.status_code == 200, response.text
            state = response.json()
        for caller, action, extra in (
            (worker, "cancel", {"reason": "First attempt"}),
            (
                lead,
                "cancellation-decision",
                {"approve": True, "comment": "First consent"},
            ),
            (
                human_resources,
                "cancellation-decision",
                {"approve": False, "comment": "Not accepted"},
            ),
            (worker, "cancel", {"reason": "Second attempt"}),
        ):
            response = await caller.post(
                f"{path}/{action}", json={"lock_version": state["lockVersion"], **extra}
            )
            assert response.status_code == 200, response.text
            state = response.json()
        assert state["cancellationManagerApproved"] is False
        denied = await human_resources.post(
            f"{path}/cancellation-decision",
            json={
                "lock_version": state["lockVersion"],
                "approve": True,
                "comment": "Attempt without fresh consent",
            },
        )
        assert denied.status_code == 409, denied.text
        assert denied.json()["error"]["code"] == "MANAGER_APPROVAL_REQUIRED"
        for caller in (lead, human_resources):
            response = await caller.post(
                f"{path}/cancellation-decision",
                json={
                    "lock_version": state["lockVersion"],
                    "approve": True,
                    "comment": "Fresh consent",
                },
            )
            assert response.status_code == 200, response.text
            state = response.json()
        assert state["status"] == "Cancelled"


@pytest.mark.asyncio
async def test_mfa_attempt_budget_survives_challenge_renewals_and_concurrency(client):
    owner, _ = await owner_client(client)
    user = await create_activated_user(owner)
    async with await spawned_client() as caller:
        await authenticate(caller, user["email"], "UserPass1!")
        setup = await caller.post("/api/v1/auth/mfa/setup")
        assert setup.status_code == 200, setup.text
        secret = setup.json()["secret"]
        assert (
            await caller.post("/api/v1/auth/mfa/confirm", json={"code": pyotp.TOTP(secret).now()})
        ).status_code == 200
        totp = pyotp.TOTP(secret)
        valid = {totp.at(datetime.now(UTC) + timedelta(seconds=offset)) for offset in (-30, 0, 30)}
        wrong = next(f"{code:06d}" for code in range(10) if f"{code:06d}" not in valid)
        tokens = []
        for _ in range(3):
            challenge = await caller.post(
                "/api/v1/auth/login",
                json={"email": user["email"], "password": "UserPass1!"},
            )
            assert challenge.status_code == 200, challenge.text
            token = challenge.json()["mfaToken"]
            tokens.append(token)
            denied = await caller.post(
                "/api/v1/auth/mfa/login", json={"token": token, "code": wrong}
            )
            assert denied.status_code == 422, denied.text
        # Two parallel guesses must consume the same account-wide persisted budget.
        outcomes = await asyncio.gather(
            *(
                caller.post("/api/v1/auth/mfa/login", json={"token": tokens[i], "code": wrong})
                for i in range(2)
            )
        )
        assert sorted(response.status_code for response in outcomes) == [422, 423]
        renewed = await caller.post(
            "/api/v1/auth/login",
            json={"email": user["email"], "password": "UserPass1!"},
        )
        assert renewed.status_code == 423, renewed.text
        for token in tokens:
            blocked = await caller.post(
                "/api/v1/auth/mfa/login", json={"token": token, "code": totp.now()}
            )
            assert blocked.status_code in (400, 423), blocked.text
            assert "nexa_session" not in blocked.cookies
    async with app.state.session_factory() as session:
        account = await session.get(User, UUID(user["id"]))
        assert account.locked_until > datetime.now(UTC)
        pending = (
            await session.scalars(
                select(OneTimeToken).where(
                    OneTimeToken.user_id == account.id,
                    OneTimeToken.purpose == "mfa_login",
                    OneTimeToken.used_at.is_(None),
                )
            )
        ).all()
        assert pending == []


@pytest.mark.parametrize("key", ["custom", "since_joining", "ytd"])
@pytest.mark.parametrize("year", [2026, 9999])
def test_history_sampling_is_bounded_and_preserves_final_endpoint(key, year):
    now = datetime(year, 9, 17, 12, tzinfo=UTC)
    start = datetime(1, 1, 1, tzinfo=UTC)
    window = PeriodWindow(key, "Whole period", start, now, start.date(), now.date())
    cutoffs = _history_cutoffs(window, now)
    assert 1 <= len(cutoffs) <= MAX_HISTORY_POINTS
    assert cutoffs[-1] == now
    assert cutoffs == sorted(set(cutoffs))
    fact = SimpleNamespace(
        id="synthetic",
        events=[],
        created_at=datetime(2026, 1, 1, tzinfo=UTC),
        terminal_at=None,
        terminal_outcome=None,
        submitted_at=datetime(2026, 2, 1, tzinfo=UTC),
        approved_at=datetime(2026, 3, 1, tzinfo=UTC),
        funded_at=datetime(2026, 4, 1, tzinfo=UTC),
    )
    history = _metric_history([fact], window, now)
    assert history["submitted"]["points"][-1]["value"] == 1
    assert history["funded"]["points"][-1]["value"] == 1
    assert history["submitted"]["points"][0]["value"] == 0
    future = datetime(year, 10, 1, tzinfo=UTC)
    assert (
        _history_cutoffs(
            PeriodWindow(key, "Future", future, future, future.date(), future.date()),
            now,
        )
        == []
    )


@pytest.mark.asyncio
async def test_legacy_patch_cannot_submit_without_submit_permission(client):
    owner, _ = await owner_client(client)
    types = (await owner.get("/api/v1/user-types")).json()["items"]
    supported = next(row for row in types if row["code"] == "BDM")
    permissions = supported["permissions"]
    changed = await owner.put(
        f"/api/v1/user-types/{supported['id']}/permissions",
        json={"permissions": ["Applications.View", "Applications.Edit"]},
    )
    assert changed.status_code == 200, changed.text
    scoped = await owner.put(
        f"/api/v1/user-types/{supported['id']}/application-scope",
        json={"application_visibility_scope": "company"},
    )
    assert scoped.status_code == 200, scoped.text
    actor = await create_activated_user(owner, user_type_code="BDM")
    await _enable_case_owner(owner, "BDM")
    bank, _, product, _ = await _catalog(owner)
    customer = await _customer(owner, "Legacy permission regression")
    application = await _create_app(
        owner,
        customer_id=customer["id"],
        bank_id=bank["id"],
        product_id=product["id"],
        case_owner_id=actor["id"],
    )
    path = f"/api/v1/applications/{application['id']}"
    async with app.state.session_factory() as session:
        before_events = list(
            await session.scalars(
                select(ApplicationEvent.id).where(
                    ApplicationEvent.application_id == UUID(application["id"])
                )
            )
        )
    async with await spawned_client() as editor:
        await authenticate(editor, actor["email"], "UserPass1!")
        denied = await editor.patch(path, json={"bank_case_number": "SECURITY-NO-SUBMIT"})
        assert denied.status_code == 403, denied.text
    async with app.state.session_factory() as session:
        retained = await session.get(Application, UUID(application["id"]))
        assert retained.submitted_at is None
        assert retained.bank_case_number is None
        after_events = list(
            await session.scalars(
                select(ApplicationEvent.id).where(ApplicationEvent.application_id == retained.id)
            )
        )
        assert set(after_events) == set(before_events)
    submitted = await owner.patch(path, json={"bank_case_number": f"SECURITY-{application['id']}"})
    assert submitted.status_code == 200, submitted.text
    assert submitted.json()["submittedAt"] is not None
    restored = await owner.put(
        f"/api/v1/user-types/{supported['id']}/permissions",
        json={"permissions": permissions},
    )
    assert restored.status_code == 200, restored.text
    restored = await owner.put(
        f"/api/v1/user-types/{supported['id']}/application-scope",
        json={"application_visibility_scope": supported["applicationVisibilityScope"]},
    )
    assert restored.status_code == 200, restored.text
    restored = await owner.put(
        f"/api/v1/user-types/{supported['id']}/case-owner",
        json={"can_be_case_owner": supported["canBeCaseOwner"]},
    )
    assert restored.status_code == 200, restored.text


@pytest.mark.asyncio
async def test_csv_formula_stage_is_literal_and_roundtrips_without_scope_bypass(client):
    owner, owner_user = await owner_client(client)
    coordinator = await create_activated_user(owner, user_type_code="COD")
    outsider = await create_activated_user(owner, user_type_code="COD")
    bank, _, product, _ = await _catalog(owner)
    customer = await _customer(owner, "CSV regression")
    application = await _create_app(
        owner,
        customer_id=customer["id"],
        bank_id=bank["id"],
        product_id=product["id"],
        case_owner_id=owner_user["id"],
    )
    async with app.state.session_factory() as session:
        row = await session.get(Application, UUID(application["id"]))
        row.routed_coordinator_id = UUID(coordinator["id"])
        stage = await session.get(WorkflowStage, row.current_stage_id)
        prior_name = stage.name
        stage.name = '  =HYPERLINK("https://example.invalid","test")'
        await session.commit()
        try:
            actor = await load_user_with_type(session, UUID(coordinator["id"]))
            payload = await current_cases_csv(session, actor)
            exported = next(
                item
                for item in csv.DictReader(StringIO(payload.decode("utf-8-sig")))
                if item["case_id"] == row.application_code
            )
            assert exported["current_stage"] == "'" + stage.name
            exported["new_stage"] = str(stage.id)
            output = StringIO(newline="")
            writer = csv.DictWriter(output, fieldnames=HEADERS)
            writer.writeheader()
            writer.writerow(exported)
            rows, errors = await validate_rows(
                session, actor, output.getvalue().encode("utf-8-sig")
            )
            assert errors == []
            assert len(rows) == 1
            foreign_actor = await load_user_with_type(session, UUID(outsider["id"]))
            _, denied = await validate_rows(
                session, foreign_actor, output.getvalue().encode("utf-8-sig")
            )
            assert denied and denied[0]["field"] == "case_id"
        finally:
            stage.name = prior_name
            await session.commit()


@pytest.mark.parametrize("activation_first", [True, False])
@pytest.mark.asyncio
async def test_kpi_edit_activation_serializes_on_real_database(
    client, monkeypatch, activation_first
):
    owner, _ = await owner_client(client)
    metric = {
        "metric_code": "funded_count",
        "weight_percent": "100",
        "direction": "higher_is_better",
    }
    created = await owner.post(
        "/api/v1/targets/kpi",
        json={"name": "Security concurrent regression", "metrics": [metric]},
    )
    assert created.status_code == 200, created.text
    card_id = created.json()["id"]
    path = f"/api/v1/targets/kpi/{card_id}"
    acquired, release = asyncio.Event(), asyncio.Event()
    original = target_service._load_scorecard
    paused = False

    async def paused_loader(session, scorecard_id, *, lock=False):
        nonlocal paused
        row = await original(session, scorecard_id, lock=lock)
        if str(scorecard_id) == card_id and lock and not paused:
            paused = True
            acquired.set()
            await asyncio.wait_for(release.wait(), 10)
        return row

    monkeypatch.setattr(target_service, "_load_scorecard", paused_loader)

    async def activate():
        return await owner.post(f"{path}/activate")

    async def edit():
        return await owner.patch(path, json={"metrics": [{**metric, "weight_percent": "10"}]})

    first = asyncio.create_task(activate() if activation_first else edit())
    second = None
    try:
        await asyncio.wait_for(acquired.wait(), 10)
        second = asyncio.create_task(edit() if activation_first else activate())
        # Assert PostgreSQL is actually waiting for the competing row lock,
        # rather than inferring overlap from task scheduling or a fixed sleep.
        for _ in range(200):
            async with app.state.engine.connect() as connection:
                waiting = await connection.scalar(
                    text(
                        "SELECT count(*) FROM pg_stat_activity WHERE datname = current_database() "
                        "AND wait_event_type = 'Lock' AND query LIKE '%kpi_scorecards%FOR UPDATE%'"
                    )
                )
            if waiting:
                break
            await asyncio.sleep(0.01)
        assert waiting, "Competing KPI transaction never waited on PostgreSQL row lock"
        assert not second.done()
    finally:
        release.set()
        responses = await asyncio.wait_for(asyncio.gather(first, *([second] if second else [])), 15)
    assert [response.status_code for response in responses] == [200, 422], [
        response.text for response in responses
    ]
    final = (await owner.get(path)).json()
    assert final["status"] == ("active" if activation_first else "draft")
    assert final["weightValid"] is activation_first
    assert Decimal(str(final["metrics"][0]["weightPercent"])) == (
        Decimal(100) if activation_first else Decimal(10)
    )


@pytest.mark.asyncio
async def test_mfa_success_resets_budget_and_challenge_cannot_replay(client):
    owner, _ = await owner_client(client)
    user = await create_activated_user(owner)
    async with await spawned_client() as caller:
        await authenticate(caller, user["email"], "UserPass1!")
        setup = await caller.post("/api/v1/auth/mfa/setup")
        assert setup.status_code == 200, setup.text
        totp = pyotp.TOTP(setup.json()["secret"])
        confirmed = await caller.post("/api/v1/auth/mfa/confirm", json={"code": totp.now()})
        assert confirmed.status_code == 200, confirmed.text
        challenge = await caller.post(
            "/api/v1/auth/login",
            json={"email": user["email"], "password": "UserPass1!"},
        )
        assert challenge.status_code == 200, challenge.text
        token = challenge.json()["mfaToken"]
        valid = {totp.at(datetime.now(UTC) + timedelta(seconds=offset)) for offset in (-30, 0, 30)}
        wrong = next(f"{code:06d}" for code in range(10) if f"{code:06d}" not in valid)
        denied = await caller.post("/api/v1/auth/mfa/login", json={"token": token, "code": wrong})
        assert denied.status_code == 422, denied.text
        async with app.state.session_factory() as session:
            account = await session.get(User, UUID(user["id"]))
            assert account.failed_login_count == 1
        outcomes = await asyncio.gather(
            *(
                caller.post("/api/v1/auth/mfa/login", json={"token": token, "code": totp.now()})
                for _ in range(2)
            )
        )
        assert sorted(response.status_code for response in outcomes) == [200, 400]
        async with app.state.session_factory() as session:
            account = await session.get(User, UUID(user["id"]))
            assert account.failed_login_count == 0
            assert account.locked_until is None
        replay = await caller.post(
            "/api/v1/auth/mfa/login", json={"token": token, "code": totp.now()}
        )
        assert replay.status_code == 400, replay.text


@pytest.mark.asyncio
async def test_kpi_concurrent_card_switches_have_one_valid_active_card(client):
    owner, _ = await owner_client(client)
    metric = {
        "metric_code": "funded_count",
        "weight_percent": "100",
        "direction": "higher_is_better",
    }
    cards = []
    for index in range(2):
        created = await owner.post(
            "/api/v1/targets/kpi",
            json={"name": f"Concurrent switch {index}", "metrics": [metric]},
        )
        assert created.status_code == 200, created.text
        cards.append(created.json()["id"])
    activated = await owner.post(f"/api/v1/targets/kpi/{cards[0]}/activate")
    assert activated.status_code == 200, activated.text
    for _ in range(3):
        outcomes = await asyncio.wait_for(
            asyncio.gather(*(owner.post(f"/api/v1/targets/kpi/{card}/activate") for card in cards)),
            10,
        )
        assert [response.status_code for response in outcomes] == [200, 200], [
            response.text for response in outcomes
        ]
        listed = await owner.get("/api/v1/targets/kpi")
        assert listed.status_code == 200, listed.text
        active = [card for card in listed.json()["items"] if card["status"] == "active"]
        assert len(active) == 1
        assert active[0]["id"] in cards
        assert active[0]["weightValid"] is True
