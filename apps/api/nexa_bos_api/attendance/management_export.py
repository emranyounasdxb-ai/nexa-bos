from __future__ import annotations

import csv
import re
from collections import defaultdict
from datetime import datetime
from html import escape
from io import BytesIO, StringIO
from pathlib import Path

from fastapi.responses import Response
from fpdf import FPDF
from fpdf.fonts import FontFace
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from nexa_bos_api.attendance.enums import BUSINESS_TZ
from nexa_bos_api.attendance.management_guards import require_storage, scoped_ids
from nexa_bos_api.attendance.management_models import AttendanceImportBatch
from nexa_bos_api.attendance.management_service import roster
from nexa_bos_api.attendance.service import _assert_filter_scope, _scoped_record_query, utcnow
from nexa_bos_api.core.exceptions import AppError
from nexa_bos_api.core.spreadsheets import spreadsheet_safe

REPORTS = {"monthly": "Monthly Attendance Register", "employee": "Employee History", "daily": "Daily Summary", "absence": "Absence", "lateness": "Lateness", "missing": "Missing Punch", "leave": "Leave", "hours": "Worked Hours", "corrections": "Corrections Audit", "imports": "Import History"}


def readable(value: str | None) -> str:
    if not value:
        return "Not recorded"
    parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    if len(value) == 10:
        return parsed.strftime("%d %b %Y")
    return parsed.astimezone(BUSINESS_TZ).strftime("%d %b %Y, %I:%M %p")


def hours(minutes: int | None) -> str:
    if minutes is None:
        return "Not recorded"
    hour, minute = divmod(minutes, 60)
    return f"{hour} hour{'s' if hour != 1 else ''} {minute} min" if hour else f"{minute} min"


async def report_data(session, actor, report, **filters):
    if report not in REPORTS:
        raise AppError(status_code=422, code="ATTENDANCE_REPORT_INVALID", message="Select an available attendance report.")
    start, end = filters["date_from"], filters["date_to"]
    if end < start or (end - start).days > 365:
        raise AppError(status_code=422, code="ATTENDANCE_RANGE_INVALID", message="Select a valid date range of at most one year.")
    if report == "employee" and not filters.get("employee_id"):
        raise AppError(status_code=422, code="ATTENDANCE_EMPLOYEE_REQUIRED", message="Select an employee for Employee History.")
    if report == "imports":
        await require_storage(session)
        allowed = await scoped_ids(session, actor)
        await _assert_filter_scope(session, allowed, employee_id=filters.get("employee_id"), office_id=filters.get("office_id"), department_id=filters.get("department_id"))
        from nexa_bos_api.attendance.management_service import employees
        staff = await employees(session, actor, filters.get("office_id"), filters.get("department_id"), filters.get("employee_id"))
        employee_ids = {str(employee.id) for employee in staff}
        batches = (await session.scalars(select(AttendanceImportBatch).order_by(AttendanceImportBatch.created_at, AttendanceImportBatch.id))).all()
        rows = []
        for batch in batches:
            if not start <= batch.created_at.astimezone(BUSINESS_TZ).date() <= end:
                continue
            scoped_rows = [row for row in batch.rows if row.get("employeeId") in employee_ids]
            if not scoped_rows:
                continue
            # Counts derive only from the current authorized/filter scope.
            rows.append({"Import Reference": batch.reference, "Filename": batch.filename, "Created": readable(batch.created_at.isoformat()), "Uploaded By": batch.uploader_snapshot.get("name") or "Not recorded", "Employee Code": batch.uploader_snapshot.get("employeeCode") or "Not recorded", "Rows": len(scoped_rows), "Valid": sum(row["classification"] == "Valid" for row in scoped_rows), "Warning": sum(row["classification"] == "Warning" for row in scoped_rows), "Error": sum(row["classification"] == "Error" for row in scoped_rows), "Status": "Imported" if batch.confirmed_at else "Staged", "Confirmed": readable(batch.confirmed_at.isoformat()) if batch.confirmed_at else "Not confirmed", "Confirmed By": (batch.confirmation_snapshot or {}).get("name") or "Not recorded"})
        return {"title": REPORTS[report], "period": {"from": start.isoformat(), "to": end.isoformat()}, "rows": rows, "columns": list(rows[0]) if rows else ["Import Reference", "Filename", "Uploaded By", "Rows", "Status"], "summary": {}}
    data = await roster(session, actor, **filters, all_rows=True, include_inactive=True)
    source = data["items"]
    if report == "absence":
        source = [row for row in source if row["recordedStatus"] == "Absent" and row["expected"]]
    elif report == "lateness":
        source = [row for row in source if row["record"] and row["record"]["isLate"]]
    elif report == "missing":
        source = [row for row in source if row["missingPunch"]]
    elif report == "leave":
        source = [row for row in source if row["approvedLeave"] or row["recordedStatus"] == "Leave"]
    if report == "daily":
        grouped = defaultdict(list)
        for row in source:
            grouped[row["attendanceDate"]].append(row)
        rows = [{"Date": readable(day), "Expected": sum(row["expected"] or 0 for row in values) if data["workingDaysConfigured"] else "Not configured", "Present": sum(row["expected"] or 0 for row in values if row["recordedStatus"] == "Present") if data["workingDaysConfigured"] else "Not configured", "Absent": sum(row["recordedStatus"] == "Absent" and bool(row["expected"]) for row in values), "Late": sum(bool(row["record"] and row["record"]["isLate"]) for row in values), "Approved Leave": sum(row["leaveWeight"] for row in values if not row["holidayWeekOff"]), "Holiday / Week Off": sum(row["holidayWeekOff"] for row in values), "Missing Punch": sum(row["missingPunch"] for row in values)} for day, values in sorted(grouped.items())]
    elif report == "corrections":
        rows = []
        for row in source:
            for correction in (row["record"] or {}).get("corrections", []):
                before, after = correction["oldValues"], correction["newValues"]
                snapshot = after.get("performedBy", {})
                rows.append({"Employee": row["employeeName"], "Employee Code": row["employeeCode"] or "Not recorded", "Office": row["office"], "Attendance Date": readable(row["attendanceDate"]), "Previous Status": before.get("status") or "Not recorded", "New Status": after.get("status") or "Not recorded", "Previous In / Out": f"{before.get('timeIn') or 'Not recorded'} / {before.get('timeOut') or 'Not recorded'}", "New In / Out": f"{after.get('timeIn') or 'Not recorded'} / {after.get('timeOut') or 'Not recorded'}", "Reason": correction["reason"], "Performed By": snapshot.get("name") or correction.get("actorName") or "Not recorded", "Staff Code": snapshot.get("employeeCode") or "Not recorded", "Designation": snapshot.get("designation") or "Not recorded", "Actor Office": snapshot.get("office") or "Not recorded", "Corrected": readable(correction["createdAt"]), "Source": after.get("source") or "Correction"})
    else:
        rows = [{"Employee": row["employeeName"], "Employee Code": row["employeeCode"] or "Not recorded", "Office": row["office"], "Date": readable(row["attendanceDate"]), "Shift": row["shift"], "Status": row["status"], "Time In": row["timeIn"] or "Not recorded", "Time Out": row["timeOut"] or "Not recorded", "Worked Hours": hours(row["workedMinutes"]), "Late Minutes": row["lateMinutes"], "Leave / Holiday": row["leaveHoliday"], "Exception": row["exception"] or "—", "Source": row["source"], "Last Updated": readable(row["lastUpdated"])} for row in source]
    return {"title": REPORTS[report], "period": {"from": start.isoformat(), "to": end.isoformat()}, "columns": list(rows[0]) if rows else ["Employee", "Employee Code", "Office", "Date", "Status"], "rows": rows, "summary": data["summary"]}


def csv_bytes(columns, rows):
    stream = StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(columns)
    for row in rows:
        writer.writerow(spreadsheet_safe(row.get(key, "")) for key in columns)
    return stream.getvalue().encode("utf-8-sig")


def colors():
    css = (Path(__file__).resolve().parents[3] / "web/app/globals.css").read_text(encoding="utf-8")
    return {key: re.search(rf"--{key}:\s*(#[0-9a-fA-F]{{6}})", css).group(1) for key in ("amafh-table-header", "amafh-table-header-text", "amafh-text", "amafh-subtle", "amafh-border")}


def report_print(data, generated):
    palette = colors()
    blocks = "".join("<tr><td><h3>" + escape(str(row.get("Employee") or row.get("Import Reference") or row.get("Date") or "Attendance")) + "</h3><dl>" + "".join(f"<div><dt>{escape(key)}</dt><dd>{escape(str(value))}</dd></div>" for key, value in row.items()) + "</dl></td></tr>" for row in data["rows"])
    variables = ";".join(f"--{key}:{value}" for key, value in palette.items())
    return f"""<!doctype html><html><head><meta charset="utf-8"><title>{escape(data['title'])}</title><style>
      :root{{{variables}}} @page{{size:A4 portrait;margin:12mm 12mm 16mm;@bottom-left{{content:'Generated {escape(generated)}';font:8px Arial;}}@bottom-right{{content:'Page ' counter(page) ' of ' counter(pages);font:8px Arial;}}}}
      body{{font:10px Arial,sans-serif;color:var(--amafh-text);margin:0;letter-spacing:normal;}} h1{{font-size:16px;margin:0 0 4px;}} header{{margin-bottom:8px;break-after:avoid;}} p{{margin:3px 0;}}
      table{{width:100%;border-collapse:separate;border-spacing:0 6px;}} thead{{display:table-header-group;}} th,h3{{background:var(--amafh-table-header);color:var(--amafh-table-header-text);padding:5px 7px;text-align:left;font-size:10px;}} h3{{margin:0;break-after:avoid;}} td{{padding:0;border:1px solid var(--amafh-border);border-radius:8px;overflow:hidden;}} tr{{break-inside:avoid;page-break-inside:avoid;}}
      dl{{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));margin:0;}} dl>div{{display:grid;grid-template-columns:80px minmax(0,1fr);gap:5px;padding:4px 7px;border-bottom:1px solid var(--amafh-border);}} dt{{font-weight:bold;}} dd{{margin:0;overflow-wrap:anywhere;}} .tools{{margin:8px 0;}} @media print{{.tools{{display:none}}}}
      </style></head><body><header><h1>AMAFH CORE — {escape(data['title'])}</h1><p>{escape(readable(data['period']['from']))} – {escape(readable(data['period']['to']))}</p><p>Generated {escape(generated)} · Asia/Dubai</p></header><div class="tools"><button onclick="window.print()">Print report</button></div><table><thead><tr><th>{escape(data['title'])}</th></tr></thead><tbody>{blocks or '<tr><td>No matching records.</td></tr>'}</tbody></table></body></html>"""


class AttendancePDF(FPDF):
    def __init__(self, title, generated, palette):
        super().__init__(orientation="P", format="A4")
        self.report_title, self.generated, self.palette = title, generated, palette
        self.report_font = "Helvetica"
        font = Path(__file__).resolve().parents[3] / "web/public/fonts/manrope/manrope-variable.ttf"
        if font.is_file():
            self.add_font("Manrope", fname=str(font), variations={"": {"wght": 400}, "B": {"wght": 600}})
            self.report_font = "Manrope"
        self.set_margins(12, 12, 12)
        self.set_auto_page_break(auto=True, margin=16)
        self.alias_nb_pages()

    def clean(self, value):
        return str(value) if self.report_font == "Manrope" else str(value).replace("—", "-").replace("–", "-").encode("latin-1", "replace").decode("latin-1")

    def header(self):
        self.set_font(self.report_font, "B", 12)
        self.multi_cell(0, 5, self.clean("AMAFH CORE — " + self.report_title), new_x="LMARGIN", new_y="NEXT")
        self.set_font(self.report_font, size=7)
        self.cell(0, 4, "Attendance report · Asia/Dubai", new_x="LMARGIN", new_y="NEXT")
        self.ln(2)

    def footer(self):
        self.set_y(-11)
        self.set_font(self.report_font, size=6)
        self.cell(0, 4, self.clean(f"Generated {self.generated} | Page {self.page_no()} of {{nb}}"), align="C")


def render_export(data, format, actor):
    generated = utcnow().astimezone(BUSINESS_TZ).strftime("%d %b %Y, %I:%M %p")
    filename = "attendance-" + data["title"].lower().replace(" ", "-")
    palette = colors()
    if format == "csv":
        content, media, extension = csv_bytes(data["columns"], data["rows"]), "text/csv; charset=utf-8", "csv"
    elif format == "print":
        content, media, extension = report_print(data, generated).encode(), "text/html; charset=utf-8", "html"
    elif format == "excel":
        book = Workbook()
        summary = book.active
        summary.title = "Summary"
        summary.append(["AMAFH CORE", data["title"]])
        for key, value in [("From", readable(data["period"]["from"])), ("To", readable(data["period"]["to"])), ("Generated", generated), ("Generated by", actor.full_name), *[(key, value if value is not None else "Not configured") for key, value in data["summary"].items()]]:
            summary.append([key, spreadsheet_safe(value)])
        rows = book.create_sheet("Attendance")
        rows.append(data["columns"])
        for row in data["rows"]:
            rows.append([spreadsheet_safe(row.get(key, "")) for key in data["columns"]])
        for sheet in book:
            sheet.sheet_view.showGridLines = False
            for line in sheet:
                sheet.row_dimensions[line[0].row].height = 32 if sheet == rows and line[0].row > 1 else 22
                for cell in line:
                    header = cell.row == 1
                    cell.fill = PatternFill("solid", fgColor=palette["amafh-table-header"].lstrip("#") if header else palette["amafh-subtle"].lstrip("#") if cell.row % 2 == 0 else "FFFFFF")
                    cell.font = Font(name="Manrope", size=10, bold=header, color="FFFFFF" if header else palette["amafh-text"].lstrip("#"))
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
                    cell.border = Border(bottom=Side(style="hair", color=palette["amafh-border"].lstrip("#")))
            for column in sheet.columns:
                letter = get_column_letter(column[0].column)
                sheet.column_dimensions[letter].width = min(38, max(16, max(len(str(cell.value or "")) for cell in column) + 2))
        from math import ceil
        for sheet in book:
            for line in sheet:
                required_lines = max(sum(max(1, ceil(len(part) / max(8, sheet.column_dimensions[get_column_letter(cell.column)].width - 2))) for part in str(cell.value or "").split("\n")) for cell in line)
                sheet.row_dimensions[line[0].row].height = min(409, max(22, required_lines * 14 + 8))
        rows.freeze_panes = "A2"
        rows.auto_filter.ref = rows.dimensions
        stream = BytesIO()
        book.save(stream)
        content, media, extension = stream.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
    else:
        pdf = AttendancePDF(data["title"], generated, palette)
        pdf.add_page()
        pdf.set_font(pdf.report_font, size=6.5)
        pdf.cell(0, 4, pdf.clean(f"{readable(data['period']['from'])} - {readable(data['period']['to'])} | Generated by {actor.full_name}"), new_x="LMARGIN", new_y="NEXT")
        rgb = tuple(int(palette["amafh-table-header"][i:i + 2], 16) for i in (1, 3, 5))
        with pdf.table(col_widths=(48, 48, 90), line_height=3.2, padding=1.3, headings_style=FontFace(emphasis="BOLD", color=(255, 255, 255), fill_color=rgb)) as table:
            heading = table.row()
            for label in ("Employee / identity", "Date / attendance", "Details / audit"):
                heading.cell(label)
            for item in data["rows"]:
                identity, attendance, details = [], [], []
                for key, value in item.items():
                    line = f"{key}: {value}"
                    target = identity if key in {"Employee", "Employee Code", "Office", "Import Reference", "Filename"} else attendance if key in {"Date", "Attendance Date", "Status", "Time In", "Time Out", "Worked Hours", "Shift", "Created", "Confirmed"} else details
                    target.append(line)
                row = table.row()
                for values in (identity, attendance, details):
                    row.cell(pdf.clean("\n".join(values) or "—"))
        content, media, extension = bytes(pdf.output()), "application/pdf", "pdf"
    return Response(content, media_type=media, headers={"Content-Disposition": f'attachment; filename="{filename}.{extension}"'})
