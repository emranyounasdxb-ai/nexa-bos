from __future__ import annotations

import csv
import re
from datetime import UTC, datetime
from html import escape
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any

from fpdf import FPDF
from fpdf.fonts import FontFace
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from nexa_bos_api.core.spreadsheets import spreadsheet_safe
from nexa_bos_api.identity.models import User


def _readable_date(value: object, precision: str = "timestamp") -> str:
    if not value:
        return "Not recorded"
    try:
        date = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        return date.strftime("%d %b %Y") if precision == "date" else date.strftime("%d %b %Y, %I:%M %p") + (" UTC" if date.utcoffset() is not None and date.utcoffset().total_seconds() == 0 else date.strftime(" %z") if date.utcoffset() is not None else "")
    except ValueError:
        return "Not recorded"


def _duration_label(seconds: object, days: object = None) -> str:
    if seconds is None:
        return f"{days:g} calendar days (Date-only record)" if isinstance(days, (int, float)) else "Not recorded"
    minutes = int(max(0, float(seconds)) // 60)
    day, remainder = divmod(minutes, 1440)
    hour, minute = divmod(remainder, 60)
    if day:
        return f"{day} days" + (f" {hour} hr" if hour else "")
    if hour:
        return f"{hour} hours" + (f" {minute} min" if minute else "")
    return f"{minute} min"


def _actor_label(actor: dict[str, Any] | None) -> str:
    actor = actor or {}
    return " · ".join(str(actor[key]) for key in ("name", "code", "designation", "office") if actor.get(key)) or "Not recorded"


def lifecycle_export_payload(history: dict[str, Any], scope: str) -> dict[str, Any]:
    """Presentation only: use existing authorized periods, totals and actor snapshots."""
    asset, lifecycle = history["asset"], history["lifecycle"]
    items, details = [], []
    for entry in lifecycle["timeline"]:
        if entry["kind"] == "event" or entry.get("excludedFromTotals"):
            details.append({"Stage": entry["title"], "Date": _readable_date(entry.get("from"), entry["precision"]),
                            "Custodian/Location": entry.get("employee") or entry.get("location") or "Not recorded",
                            "Office": entry.get("office") or "Not recorded", "Condition": entry.get("condition") or "Not recorded",
                            "Reason/Note": " · ".join(str(entry[key]) for key in ("reason", "note", "contextSource") if entry.get(key)) or "Not recorded",
                            "Performed By": _actor_label(entry.get("actor")), "Action Time": _readable_date((entry.get("actor") or {}).get("actionAt"))})
            continue
        items.append({
            "Asset Code": asset["assetCode"], "Stage": "Office/Location Transfer" if entry.get("actionCode") == "asset.office.transfer" else entry["title"],
            "Custodian Name": entry.get("employee") or "Not recorded", "Custodian Code": entry.get("employeeCode") or "Not recorded",
            "Location": entry.get("location") or "Not recorded", "Office": entry.get("office") or "Not recorded",
            "From": _readable_date(entry.get("from"), entry["precision"]),
            "To": "Current" if entry.get("active") else _readable_date(entry.get("to"), entry["precision"]),
            "Duration": _duration_label(entry.get("durationSeconds"), entry.get("durationDays")),
            "Status": entry.get("status") or "Not recorded", "Start Condition": entry.get("condition") or "Not recorded",
            "End Condition": entry.get("endCondition") or "Not recorded",
            "Reason/Note": " · ".join(str(entry[key]) for key in ("reason", "note", "endReason", "endNote", "contextSource") if entry.get(key)) or "Not recorded",
            "Opened By": _actor_label(entry.get("actor")), "Closed By": _actor_label(entry.get("endActor")),
        })
        details.append({"Stage": f"Period {len(items)} · {entry['title']}",
                        "Start Condition": entry.get("condition") or "Not recorded", "End Condition": entry.get("endCondition") or "Not recorded",
                        "Reason/Note": items[-1]["Reason/Note"],
                        "Opened By": _actor_label(entry.get("actor")), "Opened At": _readable_date((entry.get("actor") or {}).get("actionAt")),
                        "Closed By": _actor_label(entry.get("endActor")), "Closed At": _readable_date((entry.get("endActor") or {}).get("actionAt"))})
    allocation = asset.get("currentAllocation") or {}
    summary = {"Registered": _readable_date(lifecycle["registeredAt"]), "Assignments": str(lifecycle["assignments"])}
    for label, key in (("Assigned Time", "assigned"), ("Stock Time", "stock"), ("Repair Time", "repair")):
        total = lifecycle[key]
        has_recorded_period = any(e["kind"] == ("employee" if key == "assigned" else key) and e.get("durationSeconds") is not None for e in lifecycle["timeline"])
        summary[label] = "Not recorded" if not total["complete"] and not has_recorded_period else _duration_label(total["seconds"]) + (" · recorded subtotal" if not total["complete"] else "")
    if lifecycle.get("assignedDays"):
        summary["Assigned Time"] += f" + {lifecycle['assignedDays']} recorded calendar days"
    return {"title": "Asset Lifecycle History", "reportingScope": scope, "lifecycleReport": True,
            "filters": {"assetCode": asset["assetCode"], "asOf": _readable_date(lifecycle["asOf"])},
            "identity": {"Asset Code": asset["assetCode"], "Category": asset["category"]["name"],
                         "Brand/Model": " / ".join(str(asset[k]) for k in ("brand", "model") if asset.get(k)) or "Not recorded",
                         "Serial/Identity": " / ".join(str(asset[k]) for k in ("serialNumber", "imei", "iccid") if asset.get(k)) or "Not recorded",
                         "Current Status": asset["status"], "Condition": asset.get("condition") or "Not recorded",
                         "Office": (asset.get("office") or {}).get("name") or "Not recorded",
                         "Current Custodian": " · ".join(str(allocation[k]) for k in ("employeeName", "employeeCode") if allocation.get(k)) or next((str(e.get("location") or "Not recorded") for e in lifecycle["timeline"] if e.get("active") and not e.get("excludedFromTotals")), "Not recorded")},
            "summary": summary, "details": details, "items": items, "total": len(items),
            "periodActors": [(e.get("actor") or {}, e.get("endActor") or {}) for e in lifecycle["timeline"] if e["kind"] != "event" and not e.get("excludedFromTotals")], "exportPeriods": [e for e in lifecycle["timeline"] if e["kind"] != "event" and not e.get("excludedFromTotals")], "registeredTimestamp": lifecycle["registeredAt"] }


def _safe_cell(value: object) -> object:
    return spreadsheet_safe(value)


def _metadata(payload: dict[str, Any], actor: User) -> list[tuple[str, str]]:
    filters = payload.get("filters") or {}
    active_filters = (
        ", ".join(f"{key}={value}" for key, value in filters.items() if value is not None) or "None"
    )
    return [
        ("Report", str(payload.get("title") or "Asset Report")),
        ("Visibility Scope", str(payload.get("reportingScope") or "None")),
        ("Filters", active_filters),
        ("Rows", str(payload.get("total") or 0)),
        ("Generated By", f"{actor.full_name} ({actor.user_code})"),
        ("Generated At", datetime.now(UTC).strftime("%Y-%m-%d %H:%M UTC")),
    ]


def _rows(payload: dict[str, Any]) -> list[list[object]]:
    items = payload.get("items") or []
    if not items:
        if payload.get("lifecycleReport"):
            return [["Asset Code", "Stage", "Custodian Name", "Custodian Code", "Location", "Office", "From", "To", "Duration", "Status", "Start Condition", "End Condition", "Reason/Note", "Opened By", "Closed By"]]
        return [["Result"], ["No rows"]]
    headers = list(items[0].keys())
    return [headers, *[[item.get(key, "") for key in headers] for item in items]]


def build_excel(payload: dict[str, Any], actor: User) -> bytes:
    if payload.get("lifecycleReport"):
        return _build_lifecycle_excel(payload, actor)
    workbook = Workbook()
    metadata = workbook.active
    metadata.title = "Metadata"
    bold = Font(bold=True)
    for row_index, (label, value) in enumerate(_metadata(payload, actor), start=1):
        metadata.cell(row_index, 1, label).font = bold
        metadata.cell(row_index, 2, _safe_cell(value))
    results = workbook.create_sheet("Asset Report")
    for row_index, row in enumerate(_rows(payload), start=1):
        for column_index, value in enumerate(row, start=1):
            cell = results.cell(row_index, column_index, _safe_cell(value))
            if row_index == 1:
                cell.font = bold
            dimension = results.column_dimensions[get_column_letter(column_index)]
            width = min(40, max(12, len(str(value or "")) + 2))
            dimension.width = width
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _pdf_safe(value: object) -> str:
    return str(value).encode("latin-1", "replace").decode("latin-1")


def build_pdf(payload: dict[str, Any], actor: User) -> bytes:
    if payload.get("lifecycleReport"):
        return _build_lifecycle_pdf(payload, actor)
    pdf = FPDF(orientation="L")
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 8, "NEXA BOS - Asset & Inventory", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=8)
    for label, value in _metadata(payload, actor):
        pdf.cell(0, 5, _pdf_safe(f"{label}: {value}"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)
    for index, row in enumerate(_rows(payload)):
        pdf.set_font("Helvetica", "B" if index == 0 else "", 7)
        pdf.multi_cell(
            0,
            4,
            _pdf_safe(" | ".join(str(value or "") for value in row)),
            new_x="LMARGIN",
            new_y="NEXT",
        )
    return bytes(pdf.output())


def build_print_html(payload: dict[str, Any], actor: User) -> str:
    if payload.get("lifecycleReport"):
        return _build_lifecycle_print(payload, actor)
    metadata = "".join(
        f"<tr><th>{escape(label)}</th><td>{escape(value)}</td></tr>"
        for label, value in _metadata(payload, actor)
    )
    rows = _rows(payload)
    header = "".join(f"<th>{escape(str(value))}</th>" for value in rows[0])
    body = "".join(
        "<tr>" + "".join(f"<td>{escape(str(value or ''))}</td>" for value in row) + "</tr>"
        for row in rows[1:]
    )
    title = escape(str(payload.get("title") or "Asset Report"))
    return f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="utf-8"/>
  <title>{title}</title>
  <style>
    body {{ font-family: Arial, sans-serif; color: #0f172a; margin: 24px; }}
    table {{ border-collapse: collapse; width: 100%; margin-top: 16px; }}
    th, td {{ border: 1px solid #cbd5e1; padding: 6px; font-size: 11px; text-align: left; }}
    th {{ background: #f8fafc; }}
    @media print {{ button {{ display: none; }} }}
  </style>
</head>
<body>
  <button type="button" onclick="window.print()">Print</button>
  <h1>NEXA BOS - {title}</h1>
  <table>{metadata}</table>
  <table><thead><tr>{header}</tr></thead><tbody>{body}</tbody></table>
</body>
</html>
"""


def _lifecycle_report_rows(payload: dict[str, Any]) -> list[list[str]]:
    rows = [["#", "Stage", "Custodian/Location", "Office", "From", "To", "Duration", "Status", "Performed By"]]
    for index, item in enumerate(payload["items"], 1):
        custodian = " · ".join(str(item[k]) for k in ("Custodian Name", "Custodian Code") if item[k] != "Not recorded") or item["Location"]
        opened, closed = payload["periodActors"][index - 1]
        def actor_name(snapshot: dict[str, Any]) -> str:
            return " · ".join(str(snapshot[k]) for k in ("name", "code") if snapshot.get(k)) or "Not recorded"
        performed = "Opened: " + actor_name(opened) + ("; Closed: " + actor_name(closed) if closed else "")
        rows.append([str(index), item["Stage"], custodian, item["Office"], item["From"], item["To"], item["Duration"], item["Status"], performed])
    return rows


# Paper output uses the existing light-theme AMAFH semantic tokens from
# apps/web/app/globals.css. PDFs cannot resolve browser CSS variables; this is
# a snapshot of those token values, not a second palette or shared-system change.
_LIFECYCLE_PAPER_TOKENS = {
    "amafh-table-header": "#3b3d73", "amafh-table-header-text": "#ffffff",
    "amafh-border": "#e1e6f0", "amafh-text": "#192340",
    "amafh-text-secondary": "#617089", "amafh-subtle": "#f3f4fa",
}


def _paper_rgb(token: str) -> tuple[int, int, int]:
    value = _LIFECYCLE_PAPER_TOKENS[token].lstrip("#")
    return tuple(int(value[offset:offset + 2], 16) for offset in (0, 2, 4))


def _period_presentation(item: dict[str, Any]) -> list[tuple[str, str]]:
    custodian = " · ".join(str(item[key]) for key in ("Custodian Name", "Custodian Code") if item[key] != "Not recorded") or item["Location"]
    return [("Custodian / Location", custodian), ("From", item["From"]), ("To", item["To"]),
            ("Duration", item["Duration"]), ("Office", item["Office"]),
            ("Opened by", item["Opened By"]), ("Closed by", item["Closed By"])]


class _LifecyclePDF(FPDF):
    def __init__(self, payload: dict[str, Any], generated_at: str) -> None:
        super().__init__(orientation="P", format="A4")
        self.report_font = "Helvetica"
        font_path = Path(__file__).resolve().parents[3] / "web/public/fonts/manrope/manrope-variable.ttf"
        if font_path.is_file():
            self.add_font("Manrope", fname=str(font_path), variations={"": {"wght": 400}, "B": {"wght": 600}})
            self.report_font = "Manrope"
        self.report = payload
        self.generated_at = generated_at
        self.section = "Lifecycle summary"
        self.set_margins(12, 12, 12)
        self.set_auto_page_break(auto=True, margin=16)
        self.alias_nb_pages()

    def report_text(self, value: object) -> str:
        return str(value) if self.report_font == "Manrope" else _pdf_safe(value)

    def header(self) -> None:
        identity = self.report["identity"]
        self.set_text_color(*_paper_rgb("amafh-text"))
        # Requested sizes are CSS pixels; PDF font units are points (px × .75).
        self.set_font(self.report_font, "B", 16 * .75)
        self.cell(0, 6, "Asset Lifecycle History", new_x="LMARGIN", new_y="NEXT")
        self.set_font(self.report_font, size=8.5 * .75)
        for text in (f"{identity['Asset Code']} · {identity['Category']} · {identity['Brand/Model']}",
                     f"Identity: {identity['Serial/Identity']}",
                     f"{identity['Current Status']} · {identity['Current Custodian']}"):
            self.multi_cell(0, 3.2, self.report_text(text), new_x="LMARGIN", new_y="NEXT")
        self.ln(2)
        self.section_header(self.section, reserve=0)

    def footer(self) -> None:
        self.set_y(-11)
        self.set_draw_color(*_paper_rgb("amafh-border"))
        self.line(self.l_margin, self.y, self.w - self.r_margin, self.y)
        self.set_text_color(*_paper_rgb("amafh-text-secondary"))
        self.set_font(self.report_font, size=7.5 * .75)
        self.cell(self.epw * .75, 5, f"Generated {self.generated_at}")
        self.cell(self.epw * .25, 5, f"Page {self.page_no()} of {{nb}}", align="R")

    def section_header(self, title: str, reserve: float = 18) -> None:
        self.section = title
        if reserve and self.will_page_break(reserve):
            self.add_page()  # header repeats the selected section; no orphan title
            return
        self.set_fill_color(*_paper_rgb("amafh-table-header"))
        self.set_text_color(*_paper_rgb("amafh-table-header-text"))
        self.set_font(self.report_font, "B", 10.5 * .75)
        self.cell(0, 6, self.report_text(title), fill=True, new_x="LMARGIN", new_y="NEXT")
        self.set_text_color(*_paper_rgb("amafh-text"))
        self.ln(2)

    def block_metrics(self, title: str, fields: list[tuple[str, str]]) -> tuple[float, list[float], float]:
        self.set_font(self.report_font, size=8.5 * .75)
        heights = [max(3.5, float(self.multi_cell(self.epw - 43, 3.2, self.report_text(value), dry_run=True, output="HEIGHT"))) for _, value in fields]
        self.set_font(self.report_font, "B", 8.5 * .75)
        heading_height = max(6, float(self.multi_cell(self.epw - 8, 3.5, self.report_text(title), dry_run=True, output="HEIGHT")) + 3)
        return heading_height, heights, heading_height + sum(heights) + 5

    def period_block(self, title: str, fields: list[tuple[str, str]]) -> None:
        label_width, value_width = 35, self.epw - 43
        heading_height, heights, height = self.block_metrics(title, fields)
        oversized = height > self.page_break_trigger - self.t_margin - 35
        if self.will_page_break(12 if oversized else height):
            self.add_page()
        if oversized or height > self.page_break_trigger - self.y:
            # Exceptionally long legacy text flows instead of clipping or
            # creating a table row taller than a page. Section headers repeat.
            self.set_font(self.report_font, "B", 8.5 * .75)
            self.multi_cell(0, 4, self.report_text(title), new_x="LMARGIN", new_y="NEXT")
            self.set_font(self.report_font, size=8.5 * .75)
            for label, value in fields:
                self.multi_cell(0, 3.5, self.report_text(f"{label}: {value}"), new_x="LMARGIN", new_y="NEXT")
            self.ln(2)
            return
        top = self.y
        self.set_draw_color(*_paper_rgb("amafh-border"))
        self.set_line_width(.2)
        self.rect(self.l_margin, top, self.epw, height, round_corners=True, corner_radius=2)
        self.set_fill_color(*_paper_rgb("amafh-table-header"))
        self.rect(self.l_margin, top, self.epw, heading_height, style="F", round_corners=True, corner_radius=2)
        self.set_xy(self.l_margin + 4, top + 1.5)
        self.set_text_color(*_paper_rgb("amafh-table-header-text"))
        self.set_font(self.report_font, "B", 8.5 * .75)
        self.multi_cell(self.epw - 8, 3.5, self.report_text(title), new_x="LMARGIN", new_y="NEXT")
        self.set_text_color(*_paper_rgb("amafh-text"))
        y = top + heading_height + 2
        for (label, value), row_height in zip(fields, heights):
            self.set_xy(self.l_margin + 4, y)
            self.set_font(self.report_font, "B", 8 * .75)
            self.cell(label_width, 3.5, self.report_text(label))
            self.set_xy(self.l_margin + 4 + label_width, y)
            self.set_font(self.report_font, size=8.5 * .75)
            self.multi_cell(value_width, 3.2, self.report_text(value), new_x="LMARGIN", new_y="NEXT")
            y += row_height
        self.set_xy(self.l_margin, top + height + 2)


def _build_lifecycle_pdf(payload: dict[str, Any], actor: User) -> bytes:
    generated_at = datetime.now(UTC).strftime("%d %b %Y, %I:%M %p UTC")
    pdf = _LifecyclePDF(payload, generated_at)
    pdf.add_page()
    pdf.set_draw_color(*_paper_rgb("amafh-border"))
    pdf.set_text_color(*_paper_rgb("amafh-text"))
    pdf.set_font(pdf.report_font, size=8.5 * .75)
    with pdf.table(col_widths=(1, 1, 1, 1, 1), line_height=3.2, padding=1.5,
                   text_align="LEFT", headings_style=FontFace(emphasis="BOLD", color=_paper_rgb("amafh-table-header-text"), fill_color=_paper_rgb("amafh-table-header"))) as table:
        for values in (list(payload["summary"].keys()), list(payload["summary"].values())):
            row = table.row()
            for value in values:
                row.cell(pdf.report_text(value))
    pdf.ln(2)
    pdf.set_font(pdf.report_font, size=8.5 * .75)
    pdf.multi_cell(0, 3.5, pdf.report_text(f"Office: {payload['identity']['Office']} · Condition: {payload['identity']['Condition']}"), new_x="LMARGIN", new_y="NEXT")
    pdf.multi_cell(0, 3.5, pdf.report_text(f"Generated by: {actor.full_name} ({actor.user_code}) · Scope: {payload['reportingScope']}"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    first_period_height = 25.0
    if payload["items"]:
        first = payload["items"][0]
        first_period_height = pdf.block_metrics(f"1. {first['Stage']} · Status: {first['Status']}", _period_presentation(first))[2]
    pdf.section_header("Chronological lifecycle periods", reserve=min(first_period_height + 10, pdf.page_break_trigger - pdf.t_margin - 40))
    for index, item in enumerate(payload["items"], 1):
        pdf.period_block(f"{index}. {item['Stage']} · Status: {item['Status']}", _period_presentation(item))
    if not payload["items"]:
        pdf.set_font(pdf.report_font, size=8.5 * .75)
        pdf.multi_cell(0, 4, "No recorded custody periods. Available legacy actions are listed below.", new_x="LMARGIN", new_y="NEXT")
    first_detail_height = 20.0
    if payload["details"]:
        first_detail = payload["details"][0]
        first_detail_height = pdf.block_metrics(str(first_detail["Stage"]), [(label, str(value)) for label, value in first_detail.items() if label != "Stage"])[2]
    pdf.section_header("Secondary details · Individual actions and legacy records", reserve=min(first_detail_height + 10, pdf.page_break_trigger - pdf.t_margin - 40))
    for item in payload["details"]:
        pdf.period_block(str(item["Stage"]), [(label, str(value)) for label, value in item.items() if label != "Stage"])
    return bytes(pdf.output())


def _export_duration(seconds: object, days: object = None) -> str:
    """Presentation only; retain existing durations and PDF formatting."""
    if seconds is None:
        return f"{days:g} calendar {'day' if days == 1 else 'days'}" if isinstance(days, (int, float)) else "Not recorded"
    minutes = int(max(0, float(seconds)) // 60)
    day, remainder = divmod(minutes, 1440)
    hour, minute = divmod(remainder, 60)
    if day:
        return f"{day} {'day' if day == 1 else 'days'}" + (f" {hour} hr" if hour else "")
    if hour:
        return f"{hour} {'hour' if hour == 1 else 'hours'}" + (f" {minute} min" if minute else "")
    return f"{minute} min"


def _export_summary(payload: dict[str, Any]) -> dict[str, str]:
    return {label: re.sub(r"(?<![0-9])1 recorded calendar days\b", "1 calendar day", re.sub(r"(?<![0-9])1 hours\b", "1 hour", re.sub(r"(?<![0-9])1 days\b", "1 day", str(value))))
            .replace("recorded subtotal", "partial total").replace("recorded calendar days", "calendar days")
            for label, value in payload["summary"].items()}


_LIFECYCLE_EXPORT_COLUMNS = [
    "Asset Code", "Stage", "Custodian Name", "Custodian Code", "Location", "Office", "From", "To", "Is Current",
    "Duration", "Duration Minutes", "Status", "Start Condition", "End Condition", "Reason/Note",
    "Opened By Name", "Opened By Employee Code", "Opened By Designation", "Opened By Office",
    "Closed By Name", "Closed By Employee Code", "Closed By Designation", "Closed By Office",
]


def _export_period_rows(payload: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for item, period in zip(payload["items"], payload["exportPeriods"], strict=True):
        seconds, days = period.get("durationSeconds"), period.get("durationDays")
        row = {key: item[key] for key in ("Asset Code", "Stage", "Custodian Name", "Custodian Code", "Location", "Office")}
        row.update({"From": period.get("from"), "To": None if period.get("active") else period.get("to"),
                    "Is Current": bool(period.get("active")), "Duration": _export_duration(seconds, days),
                    "Duration Minutes": round(float(seconds) / 60, 6) if seconds is not None else float(days) * 1440 if days is not None else None,
                    "Status": item["Status"], "Start Condition": item["Start Condition"], "End Condition": item["End Condition"],
                    "Reason/Note": " · ".join(str(period[key]) for key in ("reason", "note", "endReason", "endNote") if period.get(key)) or "Not recorded"})
        for label, key in (("Opened By", "actor"), ("Closed By", "endActor")):
            snapshot = period.get(key) or {}
            for field, source in (("Name", "name"), ("Employee Code", "code"), ("Designation", "designation"), ("Office", "office")):
                row[f"{label} {field}"] = snapshot.get(source) or "Not recorded"
        rows.append(row)
    return rows


def _export_date(value: object) -> datetime | None:
    if not value:
        return None
    try:
        return datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except ValueError:
        return None


def _excel_date(value: object) -> datetime | None:
    date = _export_date(value)
    # Excel has no time-zone type: exact instants use explicitly labelled UTC.
    return date.astimezone(UTC).replace(tzinfo=None) if date and date.tzinfo else date


def _build_lifecycle_excel(payload: dict[str, Any], actor: User) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    lifecycle = workbook.create_sheet("Lifecycle")
    colors = {key: value.lstrip("#") for key, value in _LIFECYCLE_PAPER_TOKENS.items()}
    border = Border(bottom=Side(style="hair", color=colors["amafh-border"]))
    for sheet in (summary, lifecycle):
        sheet.sheet_view.showGridLines = False
        sheet.sheet_view.zoomScale = 85
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape" if sheet == lifecycle else "portrait"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
    summary.append(["Asset Lifecycle History", "Summary"])
    entries = {**payload["identity"], **_export_summary(payload), "Generated By": f"{actor.full_name} ({actor.user_code})",
               "Generated At (UTC)": datetime.now(UTC).replace(tzinfo=None),
               "Date/time convention": "Timestamps use UTC; date-only records retain the recorded date."}
    entries["Registered"] = _excel_date(payload.get("registeredTimestamp")) or "Not recorded"
    for label, value in entries.items():
        summary.append([label, _safe_cell(value)])
    summary.column_dimensions["A"].width = 25
    summary.column_dimensions["B"].width = 65
    summary.freeze_panes = "B2"
    lifecycle.append([key + " (UTC)" if key in ("From", "To") else key for key in _LIFECYCLE_EXPORT_COLUMNS])
    for row in _export_period_rows(payload):
        lifecycle.append([_excel_date(row[key]) if key in ("From", "To") else _safe_cell(row[key]) for key in _LIFECYCLE_EXPORT_COLUMNS])
    lifecycle.freeze_panes = "C2"
    lifecycle.auto_filter.ref = lifecycle.dimensions
    lifecycle.print_title_rows = "1:1"
    for index, key in enumerate(_LIFECYCLE_EXPORT_COLUMNS, 1):
        lifecycle.column_dimensions[get_column_letter(index)].width = (
            36 if key == "Reason/Note" else 24 if key in ("From", "To") else
            23 if "Name" in key or "Designation" in key or key == "Stage" else
            22 if "Employee Code" in key or key == "Duration Minutes" else 18)
    for sheet in (summary, lifecycle):
        for row in sheet:
            for cell in row:
                heading = cell.row == 1
                cell.font = Font(name="Manrope", size=10, color="FFFFFF" if heading else colors["amafh-text"], bold=heading or (sheet == summary and cell.column == 1))
                cell.fill = PatternFill("solid", fgColor=colors["amafh-table-header"] if heading else colors["amafh-subtle"] if cell.row % 2 == 0 else "FFFFFF")
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if isinstance(cell.value, datetime):
                    date_only = sheet == lifecycle and cell.row > 1 and cell.column in (7, 8) and payload["exportPeriods"][cell.row - 2]["precision"] == "date"
                    cell.number_format = "dd mmm yyyy" if date_only else "dd mmm yyyy, hh:mm"
                elif sheet == lifecycle and cell.column == 11:
                    cell.number_format = "0.##"
            lines = max((sum(max(1, (len(line) + int(sheet.column_dimensions[cell.column_letter].width) - 1) // int(sheet.column_dimensions[cell.column_letter].width)) for line in str(cell.value or "").split("\n")) for cell in row), default=1)
            sheet.row_dimensions[row[0].row].height = max(28 if row[0].row == 1 else 22, lines * 13 + 6)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _build_lifecycle_print(payload: dict[str, Any], actor: User) -> str:
    def safe(value: object) -> str:
        return escape(str(value))
    identity = payload["identity"]
    generated_at = datetime.now(UTC).strftime("%d %b %Y, %I:%M %p UTC")
    summary = "".join(f"<div><dt>{safe(label)}</dt><dd>{safe(value)}</dd></div>" for label, value in _export_summary(payload).items())
    periods = []
    for index, (item, period) in enumerate(zip(payload["items"], payload["exportPeriods"], strict=True), 1):
        custodian = " · ".join(str(item[key]) for key in ("Custodian Name", "Custodian Code") if item[key] != "Not recorded") or item["Location"]
        note = " · ".join(str(period[key]) for key in ("reason", "note", "endReason", "endNote") if period.get(key))
        fields = [("Custodian / Location", custodian), ("Office", item["Office"]), ("From", item["From"]), ("To", item["To"]),
                  ("Duration", _export_duration(period.get("durationSeconds"), period.get("durationDays"))),
                  ("Condition", item["Start Condition"] + " → " + item["End Condition"]),
                  ("Opened by", item["Opened By"]), ("Closed by", item["Closed By"])]
        if note:
            fields.append(("Reason / Note", note))
        facts = "".join(f"<div><dt>{safe(label)}</dt><dd>{safe(value)}</dd></div>" for label, value in fields)
        periods.append(f"<tr><td><article class='period'><h3>{index}. {safe(item['Stage'])} · {safe(item['Status'])}</h3><dl>{facts}</dl></article></td></tr>")
    tokens = ";".join(f"--{name}:{value}" for name, value in _LIFECYCLE_PAPER_TOKENS.items())
    return f"""<!DOCTYPE html><html lang="en"><head><meta charset="utf-8"><title>Asset Lifecycle History</title><style>
    :root {{ {tokens}; }}
    @font-face {{ font-family:Manrope; src:url('fonts/manrope/manrope-variable.ttf') format('truetype'); }}
    @page {{ size:A4 portrait; margin:10mm 12mm 15mm;
      @bottom-left {{ content:'Generated {generated_at}'; font:7.5px Manrope,Arial,sans-serif; color:var(--amafh-text-secondary); }}
      @bottom-right {{ content:'Page ' counter(page) ' of ' counter(pages); font:7.5px Manrope,Arial,sans-serif; }} }}
    body {{ font:8.5px Manrope,Arial,sans-serif; color:var(--amafh-text); margin:20px; letter-spacing:normal; }}
    header {{ break-inside:avoid; break-after:avoid; }} h1 {{ font-size:16px; margin:0 0 4px; }} p {{ margin:3px 0; }}
    .overview {{ break-inside:avoid; break-after:avoid; }}
    h3 {{ font-size:8.5px; margin:0; padding:4px 6px; background:var(--amafh-table-header); color:var(--amafh-table-header-text); break-after:avoid; }}
    dl {{ margin:0; display:grid; grid-template-columns:repeat(2,minmax(0,1fr)); }} dl > div {{ display:grid; grid-template-columns:80px minmax(0,1fr); padding:3px 6px; border-bottom:1px solid var(--amafh-border); break-inside:avoid; }} dt {{ font-weight:600; }} dd {{ margin:0; overflow-wrap:anywhere; }}
    .summary {{ grid-template-columns:repeat(5,minmax(0,1fr)); border:1px solid var(--amafh-border); border-top:3px solid var(--amafh-table-header); border-radius:6px; background:var(--amafh-subtle); margin:6px 0; }}
    .summary > div {{ display:block; padding:5px; border-right:1px solid var(--amafh-border); border-bottom:0; }} .summary > div:last-child {{ border-right:0; }} .summary dd {{ margin-top:3px; }}
    .periods {{ width:100%; border-collapse:separate; border-spacing:0 4px; }} .periods th {{ background:var(--amafh-table-header); color:var(--amafh-table-header-text); font-size:9px; text-align:left; padding:5px; }} .periods td {{ padding:0; }}
    thead {{ display:table-header-group; }} tr {{ break-inside:avoid; page-break-inside:avoid; }} .period {{ border:1px solid var(--amafh-border); border-radius:6px; overflow:hidden; break-inside:avoid; }}
    @media print {{ button,details,summary,nav {{ display:none!important; }} body {{ margin:0; }} .period {{ overflow:visible; }} }}
    </style></head><body><header><h1>Asset Lifecycle History</h1><p>{safe(identity['Asset Code'])} · {safe(identity['Category'])} · {safe(identity['Brand/Model'])}</p><p>Identity: {safe(identity['Serial/Identity'])}</p><p>{safe(identity['Current Status'])} · {safe(identity['Current Custodian'])}</p></header>
    <section class="overview"><dl class="summary">{summary}</dl><p>Office: {safe(identity['Office'])} · Condition: {safe(identity['Condition'])}</p><p>Prepared by: {safe(actor.full_name)} ({safe(actor.user_code)})</p></section>
    <table class="periods"><thead><tr><th>Chronological lifecycle periods · {safe(identity['Asset Code'])}</th></tr></thead><tbody>{''.join(periods) or '<tr><td>No recorded custody periods.</td></tr>'}</tbody></table></body></html>"""


def build_lifecycle_csv(payload: dict[str, Any]) -> bytes:
    output = StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow([f"{key} (UTC)" if key in ("From", "To") else key for key in _LIFECYCLE_EXPORT_COLUMNS])
    for row in _export_period_rows(payload):
        values = []
        for key in _LIFECYCLE_EXPORT_COLUMNS:
            value = row[key]
            if key in ("From", "To"):
                date = _export_date(value)
                value = date.strftime("%Y-%m-%d %H:%M:%S") if date else ""
            values.append(_safe_cell(value) if value is not None else "")
        writer.writerow(values)
    return output.getvalue().encode("utf-8-sig")
